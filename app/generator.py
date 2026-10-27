"""Generación de boletas de intimación (una por contribuyente, agrupando dominios)."""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from parser import Contribuyente

# Recursos empaquetados (solo lectura): junto al script, o en el directorio
# temporal de extracción cuando corre como ejecutable de PyInstaller.
APP_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
ASSETS_DIR = APP_DIR / "assets"
HEADER_IMG = ASSETS_DIR / "header_banner.png"

# Datos persistentes (contador de notificaciones): siempre junto al ejecutable
# o al script, nunca en el directorio temporal de extracción.
DATOS_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
DATA_DIR = DATOS_DIR / "data"
CONTADOR_PATH = DATA_DIR / "contador.json"

MESES = [
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
]

AZUL_PC = RGBColor(0x1F, 0x5C, 0x9E)
NEGRO = RGBColor(0x1A, 0x1A, 0x1A)


def _fecha_larga(dt: datetime) -> str:
    return f"Pérez, {dt.day} de {MESES[dt.month - 1]} de {dt.year}.-"


def siguiente_numero(cantidad: int) -> int:
    """Reserva y devuelve el primer número de una tanda de `cantidad` notificaciones consecutivas."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    actual = 0
    if CONTADOR_PATH.exists():
        try:
            actual = json.loads(CONTADOR_PATH.read_text(encoding="utf-8")).get("ultimo", 0)
        except (json.JSONDecodeError, OSError):
            actual = 0
    primero = actual + 1
    CONTADOR_PATH.write_text(json.dumps({"ultimo": actual + cantidad}), encoding="utf-8")
    return primero


def _p(doc: Document, texto: str = "", *, align=None, bold=False, underline=False,
       size: int | None = None, space_before: float = 0, space_after: float = 6):
    """Agrega un párrafo compacto, controlando el espaciado por estilo en vez de líneas vacías."""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if texto:
        run = p.add_run(texto)
        run.bold = bold
        run.underline = underline
        if size:
            run.font.size = Pt(size)
    return p


def _agregar_encabezado(doc: Document, numero: int):
    if HEADER_IMG.exists():
        p = _p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        run = p.add_run()
        run.add_picture(str(HEADER_IMG), width=Cm(17))

    _p(
        doc,
        f"Notificación N°: {numero:08d}",
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        bold=True,
        size=12,
        space_after=8,
    )
    _p(
        doc,
        "INTIMACIÓN - DEUDA PATENTE ÚNICA SOBRE VEHÍCULOS",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        underline=True,
        size=13,
        space_after=12,
    )


def _agregar_datos_contribuyente(doc: Document, c: Contribuyente, fecha: datetime):
    _p(doc, _fecha_larga(fecha), align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)

    for etiqueta, valor in (
        ("Nombre/Razón social", c.nombre),
        ("Domicilio", c.domicilio or "-"),
        ("CUIL/CUIT", c.cuit.replace("-", "")),
    ):
        p = _p(doc, space_after=2)
        run = p.add_run(f"{etiqueta}: ")
        run.bold = True
        p.add_run(valor)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(10)

    _p(doc, "Estimado contribuyente:", space_after=8)

    texto = (
        "De acuerdo a las constancias obrantes en nuestros registros, se le hace saber que al "
        "día de la fecha Ud. cuenta con dominio/s que registra/n deuda en "
        "concepto de Patente Única sobre Vehículos, correspondiente a los períodos fiscales "
        "comprendidos entre 1/2021 y 4/2026, a saber:"
    )
    _p(doc, texto, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)


def _agregar_tabla_dominios(doc: Document, c: Contribuyente):
    tabla = doc.add_table(rows=1, cols=3)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.style = "Light Grid Accent 1"

    encabezados = ["Dominio", "Modelo", "Tipo"]
    hdr_cells = tabla.rows[0].cells
    for i, texto in enumerate(encabezados):
        hdr_cells[i].text = ""
        run = hdr_cells[i].paragraphs[0].add_run(texto)
        run.bold = True

    for dom in sorted(c.dominios, key=lambda d: d.dominio):
        row_cells = tabla.add_row().cells
        row_cells[0].text = dom.dominio
        row_cells[1].text = dom.modelo
        row_cells[2].text = dom.tipo

    widths = [Cm(4), Cm(9), Cm(4)]
    for row in tabla.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def _agregar_pie(doc: Document):
    _p(
        doc,
        "En función a ello, se lo intima para que en el plazo perentorio de 5 (cinco) días "
        "hábiles de recibida la presente, regularice su situación de manera presencial en "
        "nuestra oficina sita en calle Sarmiento 1198 de lunes a viernes de 07:00 hs a 12:30 hs, "
        "teniendo la posibilidad de acceder a la Moratoria de Patentes vigente.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_before=8,
    )
    _p(
        doc,
        "Por cualquier duda o consulta, puede comunicarse por mail a "
        "gestiontributariaperez@gmail.com o por whatsapp al 341-2502797.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    _p(doc, "Sin más, aprovechamos la oportunidad para saludarlo atentamente.", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=40)
    _p(doc, "............................................", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _p(doc, "Oficina de Patentes Automotor", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)


def generar_boleta(c: Contribuyente, numero: int, carpeta_salida: Path, fecha: datetime | None = None) -> Path:
    fecha = fecha or datetime.now()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = NEGRO
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(6)

    _agregar_encabezado(doc, numero)
    _agregar_datos_contribuyente(doc, c, fecha)
    _agregar_tabla_dominios(doc, c)
    _agregar_pie(doc)

    carpeta_salida.mkdir(parents=True, exist_ok=True)
    cuit_limpio = c.cuit.replace("-", "")
    nombre_archivo = f"{numero:08d}_{cuit_limpio}.docx"
    destino = carpeta_salida / nombre_archivo
    doc.save(str(destino))
    return destino


def generar_boletas(
    contribuyentes: list[Contribuyente],
    carpeta_salida: Path,
    convertir_pdf: bool = False,
) -> dict:
    fecha = datetime.now()
    primero = siguiente_numero(len(contribuyentes))
    generados = []
    for i, c in enumerate(contribuyentes):
        numero = primero + i
        destino = generar_boleta(c, numero, carpeta_salida, fecha)
        generados.append(str(destino))

    resultado = {"archivos": generados, "carpeta": str(carpeta_salida), "pdf": []}

    if convertir_pdf:
        try:
            from docx2pdf import convert

            convert(str(carpeta_salida))
            resultado["pdf"] = [str(Path(p).with_suffix(".pdf")) for p in generados]
        except Exception as exc:  # noqa: BLE001
            resultado["error_pdf"] = str(exc)

    return resultado
