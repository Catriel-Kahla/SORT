"""
Lectura del archivo de deuda de Patentes (CSV o XLSX).

El archivo que exporta el sistema no trae encabezados reales: cada fila es
un dominio con datos de posición fija seguidos de bloques repetidos de
(año, período, importe original, saldo actualizado) por cada cuota adeudada.

Índices de columna (0-based) relevantes:
  1  -> dominio
  7  -> nombre / razón social
  10 -> CUIT/CUIL
  13 -> domicilio (string de ancho fijo: calle[0:25] + numero[25:30] + resto)
  14 -> modelo del vehículo
  15 -> tipo de vehículo (AUTOMOVIL, CAMION, etc.)
  18 -> localidad
  20 -> deuda total (en centavos)
  21 -> cantidad de cuotas adeudadas
  22..fin -> bloques de 4: año, período, importe original, saldo actualizado
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

IDX_DOMINIO = 1
IDX_NOMBRE = 7
IDX_CUIT = 10
IDX_DOMICILIO = 13
IDX_MODELO = 14
IDX_TIPO = 15
IDX_LOCALIDAD = 18
IDX_DEUDA_TOTAL = 20
IDX_CANT_CUOTAS = 21
IDX_BLOQUES_INICIO = 22


class ArchivoInvalido(Exception):
    pass


@dataclass
class Dominio:
    dominio: str
    modelo: str
    tipo: str
    deuda_total: float
    cant_cuotas: int


@dataclass
class Contribuyente:
    cuit: str
    nombre: str
    calle: str
    numero: str
    localidad: str
    dominios: list = field(default_factory=list)

    @property
    def domicilio(self) -> str:
        partes = [self.calle]
        if self.numero:
            partes.append(self.numero)
        return " ".join(p for p in partes if p).strip()

    @property
    def deuda_total(self) -> float:
        return sum(d.deuda_total for d in self.dominios)


def _limpiar(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _formatear_cuit(v) -> str:
    s = re.sub(r"\D", "", _limpiar(v))
    if len(s) == 11:
        return f"{s[0:2]}-{s[2:10]}-{s[10]}"
    return s or _limpiar(v)


def _parse_domicilio(raw: str) -> tuple[str, str]:
    raw = raw or ""
    calle = raw[0:25].strip()
    numero_raw = raw[25:30].strip()
    numero = numero_raw.lstrip("0") or ("0" if numero_raw else "")
    return calle, numero


def _fila_a_dominio_y_contrib(row: list) -> tuple[str, Dominio, dict]:
    dominio = _limpiar(row[IDX_DOMINIO])
    nombre = _limpiar(row[IDX_NOMBRE])
    cuit_raw = row[IDX_CUIT] if len(row) > IDX_CUIT else ""
    cuit = _formatear_cuit(cuit_raw)
    calle, numero = _parse_domicilio(_limpiar(row[IDX_DOMICILIO]) if len(row) > IDX_DOMICILIO else "")
    localidad = _limpiar(row[IDX_LOCALIDAD]) if len(row) > IDX_LOCALIDAD else ""
    modelo = _limpiar(row[IDX_MODELO]) if len(row) > IDX_MODELO else ""
    tipo = _limpiar(row[IDX_TIPO]) if len(row) > IDX_TIPO else ""
    deuda_total = _num(row[IDX_DEUDA_TOTAL]) / 100.0 if len(row) > IDX_DEUDA_TOTAL else 0.0
    cant_cuotas = int(_num(row[IDX_CANT_CUOTAS])) if len(row) > IDX_CANT_CUOTAS else 0

    dom = Dominio(
        dominio=dominio,
        modelo=modelo,
        tipo=tipo,
        deuda_total=deuda_total,
        cant_cuotas=cant_cuotas,
    )
    datos_contrib = {
        "cuit": cuit,
        "nombre": nombre,
        "calle": calle,
        "numero": numero,
        "localidad": localidad,
    }
    return cuit, dom, datos_contrib


def _iter_rows_xlsx(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    primera = next(rows_iter, None)
    if primera is None:
        return
    if _limpiar(primera[0]) == "P":
        yield list(primera)
    for row in rows_iter:
        if row is None or all(c is None for c in row):
            continue
        yield list(row)


def _iter_rows_csv(path: Path):
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        delim = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(f, delimiter=delim)
        primera = next(reader, None)
        if primera is None:
            return
        if _limpiar(primera[0]) == "P":
            yield primera
        for row in reader:
            if not row or all(not _limpiar(c) for c in row):
                continue
            yield row


def cargar_contribuyentes(path: str) -> list[Contribuyente]:
    p = Path(path)
    if not p.exists():
        raise ArchivoInvalido(f"No se encontró el archivo: {path}")

    if p.suffix.lower() in (".xlsx", ".xlsm"):
        filas = _iter_rows_xlsx(p)
    elif p.suffix.lower() in (".csv", ".txt"):
        filas = _iter_rows_csv(p)
    else:
        raise ArchivoInvalido("Formato no soportado. Usá un archivo .xlsx o .csv")

    contribuyentes: dict[str, Contribuyente] = {}
    total_filas = 0
    for row in filas:
        if row is None or len(row) <= IDX_DEUDA_TOTAL:
            continue
        dominio_val = _limpiar(row[IDX_DOMINIO])
        if not dominio_val:
            continue
        total_filas += 1
        cuit, dom, datos = _fila_a_dominio_y_contrib(row)
        if not cuit:
            continue
        c = contribuyentes.get(cuit)
        if c is None:
            c = Contribuyente(
                cuit=cuit,
                nombre=datos["nombre"],
                calle=datos["calle"],
                numero=datos["numero"],
                localidad=datos["localidad"],
            )
            contribuyentes[cuit] = c
        c.dominios.append(dom)

    if total_filas == 0:
        raise ArchivoInvalido(
            "No se encontraron filas de datos válidas en el archivo. "
            "Verificá que sea una exportación de Deuda de Patentes."
        )

    return sorted(contribuyentes.values(), key=lambda c: c.nombre.lower())
